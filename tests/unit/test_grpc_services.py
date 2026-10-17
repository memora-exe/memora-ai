"""Unit & Integration tests for all 5 gRPC AI Servicers (LibIOG3 architecture).
Uses real in-memory / local loopback gRPC channel.
"""
import io
import sys
import os
import pytest
import grpc
import fitz
import docx

# Ensure paths
ai_root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
if ai_root not in sys.path:
    sys.path.insert(0, ai_root)

proto_dir = os.path.join(ai_root, "proto")
if proto_dir not in sys.path:
    sys.path.insert(0, proto_dir)

from app.interfaces.grpc.server import create_server
from proto.ai.v1 import ai_service_pb2 as pb2
from proto.ai.v1 import ai_service_pb2_grpc as pb2_grpc


@pytest.fixture(scope="module")
def grpc_channel_and_server():
    """Start in-process ephemeral gRPC server and yield channel."""
    server = create_server(max_workers=4)
    port = server.add_insecure_port("[::]:0")
    server.start()

    channel = grpc.insecure_channel(f"localhost:{port}")
    yield channel

    channel.close()
    server.stop(grace=None)


# =====================================================================
# 1. DOCUMENT PARSER SERVICE TESTS
# =====================================================================
def test_parse_text_document(grpc_channel_and_server):
    stub = pb2_grpc.DocumentParserServiceStub(grpc_channel_and_server)
    content = (
        b"# Introduction to Memora AI\n\n"
        b"Memora is a modern knowledge management system.\n\n"
        b"It supports multi-modal document parsing and graph intelligence."
    )
    request = pb2.ParseDocumentRequest(
        file_content=content,
        file_name="intro.txt",
        mime_type="text/plain",
        options={"chunk_size": "50", "chunk_overlap": "10"},
    )
    response = stub.ParseDocument(request)

    assert "Introduction to Memora AI" in response.markdown_content
    assert response.metadata.page_count >= 1
    assert response.metadata.word_count > 5
    assert len(response.chunks) >= 1
    assert response.chunks[0].chunk_index == 0
    assert len(response.chunks[0].content) > 0


def test_parse_pdf_document(grpc_channel_and_server):
    stub = pb2_grpc.DocumentParserServiceStub(grpc_channel_and_server)

    # Create in-memory PDF with 2 pages
    pdf_doc = fitz.open()
    p1 = pdf_doc.new_page()
    p1.insert_text((50, 72), "First page content about Graph ML.")
    p2 = pdf_doc.new_page()
    p2.insert_text((50, 72), "Second page content about Dense Vector Embeddings.")
    pdf_bytes = pdf_doc.write()
    pdf_doc.close()

    request = pb2.ParseDocumentRequest(
        file_content=pdf_bytes,
        file_name="sample.pdf",
        mime_type="application/pdf",
        options={"chunk_size": "200", "chunk_overlap": "20"},
    )
    response = stub.ParseDocument(request)

    assert response.metadata.page_count == 2
    assert "First page content" in response.markdown_content
    assert "Second page content" in response.markdown_content
    assert len(response.chunks) >= 1


def test_parse_docx_document(grpc_channel_and_server):
    stub = pb2_grpc.DocumentParserServiceStub(grpc_channel_and_server)

    # Create in-memory DOCX
    doc = docx.Document()
    doc.add_heading("Docx Parsing Test", level=1)
    doc.add_paragraph("Paragraph inside word document with details.")
    docx_buf = io.BytesIO()
    doc.save(docx_buf)
    docx_bytes = docx_buf.getvalue()

    request = pb2.ParseDocumentRequest(
        file_content=docx_bytes,
        file_name="sample.docx",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        options={},
    )
    response = stub.ParseDocument(request)

    assert "Docx Parsing Test" in response.markdown_content
    assert "Paragraph inside word document" in response.markdown_content
    assert response.metadata.page_count >= 1
    assert response.metadata.word_count > 3


# =====================================================================
# 2. EMBEDDING SERVICE TESTS
# =====================================================================
def test_generate_embeddings(grpc_channel_and_server):
    stub = pb2_grpc.EmbeddingServiceStub(grpc_channel_and_server)
    texts = ["Memora knowledge graph", "Vector similarity search", "Fast stateless proxy"]
    request = pb2.GenerateEmbeddingsRequest(
        texts=texts,
        model="text-embedding-004",
        dimensions=128,
    )
    response = stub.GenerateEmbeddings(request)

    assert len(response.vectors) == 3
    for vec in response.vectors:
        assert len(vec.values) == 128
        # Ensure values are non-empty floats
        assert any(abs(v) > 0 for v in vec.values)

    assert response.usage.total_tokens > 0


def test_generate_embeddings_empty(grpc_channel_and_server):
    stub = pb2_grpc.EmbeddingServiceStub(grpc_channel_and_server)
    request = pb2.GenerateEmbeddingsRequest(texts=[])
    response = stub.GenerateEmbeddings(request)
    assert len(response.vectors) == 0


# =====================================================================
# 3. RERANK SERVICE TESTS
# =====================================================================
def test_rerank_documents(grpc_channel_and_server):
    stub = pb2_grpc.RerankServiceStub(grpc_channel_and_server)
    query = "deep learning neural networks"
    docs = [
        "Traditional French cooking techniques and recipes.",
        "Deep learning architectures using artificial neural networks.",
        "Gardening guide for planting tomatoes in spring.",
    ]
    request = pb2.RerankRequest(
        query=query,
        documents=docs,
        model="cross-encoder/ms-marco",
        top_n=2,
    )
    response = stub.Rerank(request)

    assert len(response.results) == 2
    # The deep learning document (index 1) should be ranked highest
    top_result = response.results[0]
    assert top_result.index == 1
    assert "neural networks" in top_result.document
    assert top_result.score > response.results[1].score


# =====================================================================
# 4. LLM SERVICE TESTS (UNARY + STREAMING)
# =====================================================================
def test_llm_chat_completion(grpc_channel_and_server):
    stub = pb2_grpc.LlmServiceStub(grpc_channel_and_server)
    messages = [
        pb2.LlmMessage(role="system", content="You are a helpful knowledge assistant."),
        pb2.LlmMessage(role="user", content="Summarize knowledge graph concepts."),
    ]
    request = pb2.ChatCompletionRequest(
        messages=messages,
        model="gemini-2.5-flash",
        temperature=0.5,
    )
    response = stub.ChatCompletion(request)

    assert len(response.content) > 0
    assert response.finish_reason == "stop"
    assert response.usage.prompt_tokens > 0


def test_llm_chat_completion_stream(grpc_channel_and_server):
    stub = pb2_grpc.LlmServiceStub(grpc_channel_and_server)
    messages = [
        pb2.LlmMessage(role="user", content="List three core benefits of gRPC."),
    ]
    request = pb2.ChatCompletionRequest(
        messages=messages,
        model="gemini-2.5-flash",
    )
    stream = stub.ChatCompletionStream(request)

    collected_content = []
    has_final_chunk = False

    for chunk in stream:
        if chunk.delta_content:
            collected_content.append(chunk.delta_content)
        if chunk.is_done:
            has_final_chunk = True
            assert chunk.usage.total_tokens > 0

    assert has_final_chunk is True
    full_text = "".join(collected_content)
    assert len(full_text) > 0


# =====================================================================
# 5. GRAPH ML SERVICE TESTS
# =====================================================================
def test_detect_communities(grpc_channel_and_server):
    stub = pb2_grpc.GraphMlServiceStub(grpc_channel_and_server)
    # Community A: n1, n2, n3
    # Community B: n4, n5, n6
    nodes = ["n1", "n2", "n3", "n4", "n5", "n6"]
    edges = [
        pb2.GraphEdgeInput(source_id="n1", target_id="n2", weight=1.0),
        pb2.GraphEdgeInput(source_id="n2", target_id="n3", weight=1.0),
        pb2.GraphEdgeInput(source_id="n1", target_id="n3", weight=1.0),
        pb2.GraphEdgeInput(source_id="n4", target_id="n5", weight=1.0),
        pb2.GraphEdgeInput(source_id="n5", target_id="n6", weight=1.0),
        pb2.GraphEdgeInput(source_id="n4", target_id="n6", weight=1.0),
    ]
    request = pb2.DetectCommunitiesRequest(
        node_ids=nodes,
        edges=edges,
        algorithm="louvain",
    )
    response = stub.DetectCommunities(request)

    assert len(response.clusters) == 2
    cluster_node_sets = [set(c.node_ids) for c in response.clusters]
    assert {"n1", "n2", "n3"} in cluster_node_sets
    assert {"n4", "n5", "n6"} in cluster_node_sets


def test_detect_communities_empty(grpc_channel_and_server):
    stub = pb2_grpc.GraphMlServiceStub(grpc_channel_and_server)
    request = pb2.DetectCommunitiesRequest(node_ids=[], edges=[])
    response = stub.DetectCommunities(request)
    assert len(response.clusters) == 0


# =====================================================================
# 6. LIBIOG3 STRUCTURE VALIDATION TESTS
# =====================================================================
def test_libiog3_adapter_structure():
    """Verify adapters follow LibIOG3 naming and backward compatibility wrappers."""
    from app.infrastructure.adapters.embedder_adapter import EmbedderAdapter, EmbeddingAdapter
    from app.infrastructure.adapters.reranker_adapter import RerankerAdapter, RerankAdapter
    from app.infrastructure.adapters.embedding_adapter import EmbeddingAdapter as WrapperEmbeddingAdapter
    from app.infrastructure.adapters.rerank_adapter import RerankAdapter as WrapperRerankAdapter
    from app.infrastructure.adapters import (
        DocumentParserAdapter,
        EmbedderAdapter as ExportedEmbedder,
        RerankerAdapter as ExportedReranker,
        LlmAdapter,
        GraphMlAdapter,
    )

    assert EmbedderAdapter is EmbeddingAdapter
    assert EmbedderAdapter is WrapperEmbeddingAdapter
    assert EmbedderAdapter is ExportedEmbedder

    assert RerankerAdapter is RerankAdapter
    assert RerankerAdapter is WrapperRerankAdapter
    assert RerankerAdapter is ExportedReranker

    embedder = EmbedderAdapter()
    res = embedder.generate_embeddings(["test"])
    assert len(res.vectors) == 1

    reranker = RerankerAdapter()
    ranked = reranker.rerank("test", ["test doc", "other"])
    assert len(ranked) == 2


def test_libiog3_proto_structure():
    """Verify proto directory and stubs follow LibIOG3 structure."""
    from proto.ai.v1 import ai_service_pb2 as proto_pb2
    from proto.ai.v1 import ai_service_pb2_grpc as proto_pb2_grpc
    import proto.ai_service_pb2 as root_pb2
    import proto.ai_service_pb2_grpc as root_pb2_grpc

    assert hasattr(proto_pb2, "ParseDocumentRequest")
    assert hasattr(proto_pb2, "GenerateEmbeddingsRequest")
    assert hasattr(proto_pb2, "RerankRequest")
    assert hasattr(proto_pb2, "ChatCompletionRequest")
    assert hasattr(proto_pb2, "DetectCommunitiesRequest")

    assert hasattr(proto_pb2_grpc, "DocumentParserServiceStub")
    assert hasattr(proto_pb2_grpc, "EmbeddingServiceStub")
    assert hasattr(proto_pb2_grpc, "RerankServiceStub")
    assert hasattr(proto_pb2_grpc, "LlmServiceStub")
    assert hasattr(proto_pb2_grpc, "GraphMlServiceStub")

    assert hasattr(root_pb2, "ParseDocumentRequest")
    assert hasattr(root_pb2_grpc, "DocumentParserServiceStub")

